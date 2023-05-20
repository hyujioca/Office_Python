from docx import Document
document = Document("sample.docx")
total_characters = 0
for paragraph in document.paragraphs:
    total_characters += len(paragraph.text)
print(total_characters)
document.add_picture("download.jpg")
document.save('sample_answer.docx')